#!/usr/bin/env python3
"""Extract text and tables from a .docx file.

Usage:
    uv run --with python-docx python3 docx-extraction.py <file.docx>
    # Or if python-docx is already installed:
    python3 docx-extraction.py <file.docx>

Handles Chinese-encoded filenames gracefully: pass the path as argv[1], or
modify the glob in `main()` to auto-discover files by pattern.
"""
import sys
import glob
from pathlib import Path
from docx import Document


def extract_text(doc: Document) -> str:
    """Extract all paragraph text, skip empty lines."""
    lines = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            lines.append(text)
    return "\n".join(lines)


def extract_tables(doc: Document) -> list[str]:
    """Extract all tables as pipe-separated text."""
    results = []
    for i, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(" | ".join(cells))
        results.append(f"\n--- 表格 {i+1} ---\n" + "\n".join(rows))
    return results


def main():
    if len(sys.argv) > 1:
        doc_path = Path(sys.argv[1])
    else:
        # Auto-discover: modify the glob pattern as needed
        files = glob.glob("*.docx")
        if not files:
            print("No .docx file provided and none found in cwd.", file=sys.stderr)
            sys.exit(1)
        doc_path = Path(files[0])

    if not doc_path.exists():
        print(f"File not found: {doc_path}", file=sys.stderr)
        sys.exit(1)

    print(f"Reading: {doc_path}")
    print("=" * 60)

    doc = Document(str(doc_path))

    # Paragraphs
    text = extract_text(doc)
    if text:
        print(text)

    # Tables
    for table_text in extract_tables(doc):
        print(table_text)

    print(f"\n--- 统计 ---")
    print(f"段落数: {len(doc.paragraphs)}")
    print(f"表格数: {len(doc.tables)}")


if __name__ == "__main__":
    main()
