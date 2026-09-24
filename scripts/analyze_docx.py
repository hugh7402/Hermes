# -*- coding: utf-8 -*-
"""分析6月docx的格式细节"""
from docx import Document

d = Document('/opt/data/cache/documents/doc_86e25508d2ed_团队月度重点项目汇报材料_2026年6月_sl.docx')

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

print('=== 段落列表 ===')
for i, p in enumerate(d.paragraphs):
    if not p.text.strip():
        continue
    style = p.style.name if p.style else 'None'
    runs_info = []
    for r in p.runs[:2]:
        f = r.font
        ea = None
        if r._element.rPr is not None and r._element.rPr.rFonts is not None:
            ea = r._element.rPr.rFonts.get(W + 'eastAsia')
        runs_info.append(f'name={f.name},ea={ea},size={f.size},bold={f.bold}')
    print(f'[{i}] style={style} | {p.text[:45]} | {runs_info}')

print()
print('=== 对齐/缩进（前20个非空段） ===')
cnt = 0
for i, p in enumerate(d.paragraphs):
    if not p.text.strip():
        continue
    pf = p.paragraph_format
    print(f'[{i}] align={pf.alignment} first_indent={pf.first_line_indent} line_spacing={pf.line_spacing} space_before={pf.space_before} space_after={pf.space_after} | {p.text[:30]}')
    cnt += 1
    if cnt >= 20:
        break
