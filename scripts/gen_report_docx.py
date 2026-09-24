# -*- coding: utf-8 -*-
"""以6月docx为模板生成7-8月汇报材料docx，格式保持一致"""
import shutil, copy
from docx import Document

SRC = '/opt/data/cache/documents/doc_86e25508d2ed_团队月度重点项目汇报材料_2026年6月_sl.docx'
DST = '/opt/data/OutPut Box/团队月度重点项目汇报材料_2026年7月-8月/团队月度重点项目汇报材料_2026年7月-8月.docx'

shutil.copy(SRC, DST)
d = Document(DST)

# 收集关键段落元素
paras = d.paragraphs
title_elm = copy.deepcopy(paras[0]._element)   # 标题 18pt bold
info1_elm = copy.deepcopy(paras[1]._element)   # 团队名称 10pt
info2_elm = copy.deepcopy(paras[2]._element)   # 汇报时间 10pt
h2_elm    = copy.deepcopy(paras[4]._element)   # 章节标题 Heading2
body_elm  = copy.deepcopy(paras[6]._element)   # 正文段（首行缩进）
num_elm   = copy.deepcopy(paras[12]._element)  # 编号段（无缩进）

def set_text(elm, text):
    """清空段落文本，写入第一个run"""
    from docx.oxml.ns import qn
    runs = elm.findall(qn('w:r'))
    for r in runs[1:]:
        elm.remove(r)
    if runs:
        t = runs[0].find(qn('w:t'))
        if t is None:
            t = runs[0].makeelement(qn('w:t'), {})
            runs[0].append(t)
        t.text = text
    else:
        r = elm.makeelement(qn('w:r'), {})
        t = r.makeelement(qn('w:t'), {})
        t.text = text
        r.append(t)
        elm.append(r)

# 内容结构：('h', 章节标题) / ('body', 正文段) / ('num', 编号段)
content = [
    ('title', '团队月度重点项目汇报材料'),
    ('info1', '团队名称：联通数智-智慧民政项目组   汇报人：          '),
    ('info2', '汇报时间：2026年7月-8月'),
    ('h', '一、重点业务产品打造情况'),
    ('num', '1. 统一UI标准落地。7月20日专班确定统一UI原型设计并启动标准制定，7月30日完成效果图及高保真设计稿，标准正式下发，7月31日召开宣贯会组织各系统按统一标准推进开发。'),
    ('num', '2. "陕西民政通"小程序改版。围绕智能问答升级完成整体重构方案，与省数政局协调落实服务器资源和服务接口下发，问答智能体本地开发推进中，预计8月底前完成改版及智能问答演示版本。'),
    ('num', '3. 业务看板产品化。按省民政领域技能大赛展示要求，完成13个业务看板UI设计及前端开发，涵盖养老、救助、儿童福利、慈善等业务领域，具备数据汇聚、态势感知、决策辅助能力。'),
    ('h', '二、重点项目攻坚与落地成果'),
    ('num', '1. 需求确认集中收官。7月15日低收入人口信息监测系统核对域原型、福彩系统原型完成确认；7月20日统一UI原型确定；7月24日儿童福利系统原型完成确认，十二个业务系统需求确认全部完成。8月区划地名系统工作台和业务看板原型完成确认并签署确认单，统一区划完成最终版本输出，形成全省民政统一行政区划底座。'),
    ('num', '2. 系统开发多点提速。截至7月底6个系统正式上线（养老、救助、殡葬、儿童、社会组织、民政大脑），3个试运行，4个开发中。低收入人口信息监测系统完成政府购买服务、救助移动端系统开发；残疾人福利系统第二批（精康、民康、公共服务）开发完成；殡葬系统完成统计分析报表、红白理事会管理功能；社会组织党建网站开发完成；项目资金监管系统完成专项资金功能。'),
    ('num', '3. 8月关键节点突破。福彩系统完成投注机申请及合同签署全流程端到端测试并确认功能；慈善综合管理系统部署至生产环境；民政电子档案完成部分处室对接清单标准制定；公共支撑完成物联网对接标准、统一API和统一CLI对接注册。'),
    ('num', '4. 数据治理深化。养老系统完成民政部122张表整体数据回流，累计更新50万人能力数据、新增2.3万人；婚姻系统与档案系统实现"一键归档→预归档→档案系统接收→四性监测→人工审核→装盒归档"全流程闭环；数据安全系统完成高性能云环境迁移适配。'),
    ('num', '5. 平台运行稳中有升。截至7月底平台累计开通账号69,820个，累计登录535.8万次（当月新增34.12万），累计点击量6,029.3万次；对外共享数据6,132.70万条、归集厅外数据9.6亿条，9个业务板块实现与民政部数据对接；累计组织基层培训49次，处理账号配置申请1,712次，解答业务咨询2,717次。'),
    ('h', '三、商机拓展与储备转化情况'),
    ('num', '1. 农资监管平台调研。7月27日参加农业三资监管平台调研，28日完成调研报告编写，30日向厅里汇报调研报告和项目资金监管情况，为涉农领域信息化拓展储备方向。'),
    ('num', '2. 村（社区）系统建设。8月5日组织村组织审核登记系统讨论，8月6日完成村社区系统建设建议书，推动基层治理领域新业务落地。'),
    ('num', '3. 外部合作对接。与蚂蚁集团沟通养老优待证办理、民政通合作及养老领域参与内容；对标调研医保、住建一体化平台；8月7日对接捐赠系统情况，拓展慈善领域业务空间。'),
    ('h', '四、现存问题、卡点与风险事项'),
    ('num', '1. 处室原型确认持续滞后。8月15项延期任务根因是处室原型确认滞后，社会力量帮扶、低收入人口、养老、区划地名、殡葬等系统原型待处室或厅领导确认，部分任务自7月延续至今。'),
    ('num', '2. 档案归档标准长期未统一。民政电子档案与婚姻、救助、慈善、社会组织、养老等板块对接因归档标准及业务方案未明确持续延期，其中婚姻系统需求确认自6月至今未获实质进展。'),
    ('num', '3. 政策文件和指标数据收集滞后。民政通智能问答所需政策文件仅婚姻、老龄处室提供，业务看板展示指标数据普遍未提供，直接影响技能大赛演示进度。'),
    ('num', '4. 平台整体进度偏慢。截至8月23日246项任务累计完成109项（进度44%），延期38项，当期任务完成率20%，后续销号压力较大。'),
    ('h', '五、工作计划与攻坚举措'),
    ('num', '1. 9月初展示攻坚。围绕省民政领域技能大赛智慧民政展示，完成民政通智能问答演示版本和13个业务看板数据填充，推动各处室8月25日前提交政策文件和指标数据。'),
    ('num', '2. 延期任务销号。第七双周（8月24日-9月6日）滚动推进80项任务，聚焦处室原型确认提速、档案标准协调落地、政策文件及指标数据收集、外部接口收尾四项攻坚方向。'),
    ('num', '3. 民政通改版上线。8月底前完成小程序改版及智能问答演示版本，推进语音问答、智能导办功能开发上线。'),
    ('num', '4. 系统上线推广。推动福彩、婚姻、区划地名、慈善等系统完成开发和试运行，组织基层业务人员集中培训，加快推广覆盖。'),
    ('h', '六、其他事项'),
]

body = d.element.body
# 保留 sectPr
sectPr = body.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr')
# 删除所有段落
for p in list(body.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p')):
    body.remove(p)

# 重建
for kind, text in content:
    if kind == 'title':
        elm = copy.deepcopy(title_elm)
    elif kind == 'info1':
        elm = copy.deepcopy(info1_elm)
    elif kind == 'info2':
        elm = copy.deepcopy(info2_elm)
    elif kind == 'h':
        elm = copy.deepcopy(h2_elm)
    elif kind == 'body':
        elm = copy.deepcopy(body_elm)
    else:
        elm = copy.deepcopy(num_elm)
    set_text(elm, text)
    body.insert(list(body).index(sectPr) if sectPr is not None else len(body), elm)

d.save(DST)
print('saved:', DST)

# 验证：重新打开，统计段落数和字体
d2 = Document(DST)
print('paragraphs:', len(d2.paragraphs))
for p in d2.paragraphs[:6]:
    runs = p.runs[:1]
    f = runs[0].font if runs else None
    sz = f.size if f else None
    print(f'[{p.style.name}] size={sz} bold={f.bold if f else None} | {p.text[:40]}')
