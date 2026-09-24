"""
微信文档自动入库脚本
- 扫描 Hermes 文档缓存目录 /opt/data/cache/documents/
- 将你通过微信发给我的文件自动转换为 Markdown
- 入库到 Obsidian Vault 的 01-WeiXin/ 目录
- 调用 note_enhance.py 进行 AI 增强（摘要+标签+关联）
"""
import subprocess, sys, os, re, shutil, hashlib
from pathlib import Path
from datetime import datetime

CACHE_DIR = "/opt/data/cache/documents"
WEIXIN_DIR = "/opt/data/Obsidian Vault/Obsidian Vault/01-WeiXin"
VAULT = "/opt/data/Obsidian Vault/Obsidian Vault"
RAW = f"{VAULT}/raw"
PROCESSED_LOG = f"{VAULT}/.ingested_weixin"
MAX_PER_RUN = 20


def get_processed():
    """Load set of already-processed file hashes"""
    processed = set()
    if os.path.exists(PROCESSED_LOG):
        with open(PROCESSED_LOG) as f:
            for line in f:
                processed.add(line.strip())
    return processed


def mark_processed(filepath):
    """Record file as processed"""
    h = hashlib.md5(Path(filepath).read_bytes()).hexdigest()
    with open(PROCESSED_LOG, 'a') as f:
        f.write(f"{h}\n")
    return h


def is_already_processed(filepath):
    h = hashlib.md5(Path(filepath).read_bytes()).hexdigest()
    return h in get_processed()


def get_source_from_filename(filename):
    """Extract a readable source label from cached filename.
    Cached files look like: doc_{uuid12}_{original_name.ext}"""
    # Strip the uuid prefix
    name = Path(filename).name
    if name.startswith("doc_") and len(name) > 17:
        name = name[17:]  # skip "doc_" + 12 hex chars + "_"
    return name


def convert_to_markdown(filepath):
    """Convert document to markdown text"""
    ext = Path(filepath).suffix.lower()
    print(f"  转换 {Path(filepath).name} ({ext})...")

    try:
        if ext == '.docx':
            result = subprocess.run(
                ['uv', 'run', '--with', 'python-docx', 'python3', '-c', f'''
import docx
doc = docx.Document("{filepath}")
text = []
for p in doc.paragraphs:
    if p.style.name.startswith("Heading"):
        level = int(p.style.name.split()[-1]) if p.style.name.split()[-1].isdigit() else 1
        text.append("#" * level + " " + p.text)
    else:
        text.append(p.text)
print("\\\\n".join(text))
'''],
                capture_output=True, text=True, timeout=30
            )
            return result.stdout.strip()

        elif ext == '.pdf':
            # 先尝试 pymupdf 提取文字
            result = subprocess.run(
                ['uv', 'run', '--with', 'pymupdf', 'python3', '-c', f'''
import fitz
doc = fitz.open("{filepath}")
text = []
for page in doc:
    text.append(page.get_text())
print("\\\\n".join(text))
'''],
                capture_output=True, text=True, timeout=30
            )
            text = result.stdout.strip()
            # 扫描件 → OCR
            if len(text) < 50:
                print(f"  📸 扫描件，切换 OCR...")
                ocr_result = subprocess.run(
                    ['/opt/hermes/.venv/bin/python3', '/opt/data/pdf_ocr.py', filepath],
                    capture_output=True, text=True, timeout=400
                )
                if ocr_result.returncode == 0 and ocr_result.stdout.strip():
                    text = ocr_result.stdout.strip()
                else:
                    print(f"  ⚠️ OCR 失败: {ocr_result.stderr[:200]}")
            return text

        elif ext == '.pptx':
            result = subprocess.run(
                ['uv', 'run', '--with', 'python-pptx', 'python3', '-c', f'''
from pptx import Presentation
prs = Presentation("{filepath}")
text = []
for slide in prs.slides:
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text.append(para.text)
print("\\\\n".join(text))
'''],
                capture_output=True, text=True, timeout=30
            )
            return result.stdout.strip()

        elif ext in ['.xls', '.xlsx']:
            result = subprocess.run(
                ['uv', 'run', '--with', 'openpyxl', '--with', 'xlrd', 'python3', '-c', f'''
import sys
path = "{filepath}"
text = []

try:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        text.append(f"## {{sheet_name}}")
        for row in ws.iter_rows(values_only=True):
            row_text = " | ".join([str(c) if c is not None else "" for c in row])
            if row_text.strip():
                text.append(row_text)
        text.append("")
except Exception:
    import xlrd
    wb = xlrd.open_workbook(path)
    for sheet in wb.sheets():
        text.append(f"## {{sheet.name}}")
        for row_idx in range(min(sheet.nrows, 200)):
            row_text = " | ".join([str(sheet.cell_value(row_idx, c)) for c in range(sheet.ncols)])
            if row_text.strip():
                text.append(row_text)
        text.append("")

print("\\\\n".join(text))
'''],
                capture_output=True, text=True, timeout=30
            )
            return result.stdout.strip()

        elif ext == '.txt':
            try:
                return Path(filepath).read_text(encoding='utf-8').strip()
            except UnicodeDecodeError:
                # Try gbk for Chinese text files
                try:
                    return Path(filepath).read_text(encoding='gbk').strip()
                except Exception:
                    return Path(filepath).read_text(encoding='latin-1').strip()

        elif ext in ['.md', '.markdown']:
            return Path(filepath).read_text(encoding='utf-8').strip()

        else:
            print(f"  ⚠️ 不支持的格式: {ext}")
            return None

    except Exception as e:
        print(f"  ❌ 转换失败: {e}")
        return None


def extract_title(text):
    """Extract title from first # heading"""
    for line in text.split('\n'):
        if line.startswith('# '):
            return line[2:].strip()
    return None


def clean_title(title):
    """Clean filename to use as note title"""
    # Remove extension
    name = title.rsplit('.', 1)[0]
    # Remove uuid prefix if present
    if name.startswith("doc_") and len(name) > 17:
        name = name[17:]
    # Clean special chars
    name = name.replace(' ', '-').replace('/', '-')[:80]
    return name


def save_note(filename, text, source_label=""):
    """Save markdown to 01-WeiXin/"""
    safe_name = clean_title(filename) + '.md'
    path = os.path.join(WEIXIN_DIR, safe_name)

    # Add source metadata
    source_line = f"> **来源**：微信 — {source_label or filename}"

    # Ensure there's a # Title
    has_h1 = any(line.startswith('# ') for line in text.split('\n'))
    base_title = filename.rsplit('.', 1)[0]
    # Clean title for display
    display_title = base_title
    if display_title.startswith("doc_") and len(display_title) > 17:
        display_title = display_title[17:]

    if not has_h1:
        text = f"# {display_title}\n\n{source_line}\n\n---\n\n{text}"
    else:
        # Insert source after first heading
        lines = text.split('\n')
        for i, line in enumerate(lines):
            if line.startswith('# '):
                lines.insert(i+1, '')
                lines.insert(i+2, source_line)
                lines.insert(i+3, '')
                lines.insert(i+4, '---')
                break
        text = '\n'.join(lines)

    Path(path).write_text(text, encoding='utf-8')
    return path


def update_index(notes, source="微信"):
    """Update index.md with new 01-WeiXin notes"""
    index_path = os.path.join(VAULT, 'index.md')
    content = Path(index_path).read_text(encoding='utf-8')

    for title, filename in notes:
        entry = f"- [[{title}]] — 微信自动摄入"
        if entry not in content:
            marker = '## 实体'
            content = content.replace(marker, f'{entry}\n\n{marker}')

    # Update total count for 01-WeiXin
    wx_count = len(list(Path(WEIXIN_DIR).glob('*.md')))
    content = re.sub(r'共 \d+ 页', f'共 {wx_count + 584 + 2} 页', content)
    content = re.sub(r'更新：\d{4}-\d{2}-\d{2}', f'更新：{datetime.now().strftime("%Y-%m-%d")}', content)

    Path(index_path).write_text(content, encoding='utf-8')


def update_log(files, source="微信"):
    """Append to log.md"""
    log_path = os.path.join(VAULT, 'log.md')
    today = datetime.now().strftime('%Y-%m-%d')
    entry = f"\n## [{today}] weixin_ingest | 微信文档入库 {len(files)} 篇\n"
    for f in files:
        entry += f"- {Path(f).name}\n"

    with open(log_path, 'a') as f:
        f.write(entry)


def main():
    # Ensure target dir exists
    Path(WEIXIN_DIR).mkdir(parents=True, exist_ok=True)

    if not os.path.exists(CACHE_DIR):
        print(f"❌ 缓存目录不存在: {CACHE_DIR}")
        return

    files = sorted(Path(CACHE_DIR).iterdir())
    # Filter: only document-like files, skip empty ones
    valid_exts = {'.docx', '.pdf', '.pptx', '.xls', '.xlsx', '.txt', '.md', '.markdown'}
    docs = [f for f in files if f.is_file() and f.suffix.lower() in valid_exts and f.stat().st_size > 0]

    if not docs:
        print(f"📭 缓存目录无文档文件 ({len(files)} 个文件但无可识别格式)")
        return

    unprocessed = [f for f in docs if not is_already_processed(str(f))]
    if not unprocessed:
        print(f"📭 无新微信文档（缓存中共 {len(docs)} 个，均已处理过）")
        return

    print(f"📊 微信缓存共 {len(unprocessed)} 个待处理文件，本次处理 {min(len(unprocessed), MAX_PER_RUN)} 个")

    new_notes = []

    for f in unprocessed[:MAX_PER_RUN]:
        print(f"\n📄 {f.name}")
        text = convert_to_markdown(str(f))
        if not text:
            print(f"  ⤵️  保留文件，下次再试")
            continue

        source_label = get_source_from_filename(f.name)
        safe_name = clean_title(f.name) + '.md'
        path = save_note(safe_name, text, source_label)

        # Run note_enhance on the new note (传绝对路径)
        print(f"  🤖 增强...")
        subprocess.run(
            ['python3', '/opt/data/note_enhance.py', path],
            timeout=300
        )

        mark_processed(str(f))

        title = extract_title(text)
        if title:
            new_notes.append((title, safe_name))
        print(f"  ✅ 已入库: 01-WeiXin/{Path(path).name}")

    if new_notes:
        update_index(new_notes)
        update_log([n[1] for n in new_notes])
        print(f"\n✨ 微信文档入库 {len(new_notes)} 篇新笔记")
    else:
        print("\n📭 无新文档需处理")


if __name__ == '__main__':
    main()
